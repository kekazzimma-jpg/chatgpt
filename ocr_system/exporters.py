from __future__ import annotations
import base64
import json
import os
import shutil
from pathlib import Path
from typing import Any, Dict, List
import fitz
import html
from document_model import spans_look_incoherent, spans_majority_style
def export_json(document: Dict[str, Any], out_path: Path) -> Path:
    out_path.write_text(json.dumps(document, ensure_ascii=False, indent=2), encoding="utf-8")
    return out_path
def _style_wrap_md(text: str, span: Dict[str, Any]) -> str:
    t = text
    if span.get("all_caps"):
        t = t.upper()
    if span.get("underline"):
        t = f"<u>{t}</u>"
    if span.get("italic"):
        t = f"*{t}*"
    if span.get("bold"):
        t = f"**{t}**"
    return t
def _block_text_md(block: Dict[str, Any]) -> str:
    content = str(block.get("content", ""))
    spans = block.get("inline_spans", []) if isinstance(block.get("inline_spans"), list) else []
    if spans:
        # Safety net: se gli spans sembrano incoerenti col content (dovrebbero essere già
        # stati normalizzati, ma la rete non fa male), collassiamo in un unico span piatto
        # conservando però i flag di stile dominanti — non rendiamo tutto neutro.
        if spans_look_incoherent(spans, content):
            majority = spans_majority_style(spans)
            return _style_wrap_md(content, majority)
        return "".join(_style_wrap_md(str(s.get("text", "")), s) for s in spans)
    return content
def _fallback_list_items_from_content(content: str) -> List[str]:
    lines = [ln.strip() for ln in content.splitlines() if ln.strip()]
    out: List[str] = []
    for ln in lines:
        ln2 = ln
        if ln.startswith("- "):
            ln2 = ln[2:].strip()
        out.append(ln2)
    return out
def _render_inline_html(spans: List[Dict[str, Any]], fallback: str) -> str:
    if not spans:
        return html.escape(fallback)
    out = []
    for sp in spans:
        txt = html.escape(str(sp.get("text", "")))
        if sp.get("all_caps"):
            txt = txt.upper()
        if sp.get("underline"):
            txt = f"<u>{txt}</u>"
        if sp.get("italic"):
            txt = f"<em>{txt}</em>"
        if sp.get("bold"):
            txt = f"<strong>{txt}</strong>"
        out.append(txt)
    return "".join(out) if out else html.escape(fallback)
def _align_css(style: Dict[str, Any]) -> str:
    align = str(style.get("alignment", "left")).lower()
    if align not in {"left", "center", "right", "justify"}:
        align = "left"
    return f"text-align:{align};"
def export_markdown(document: Dict[str, Any], out_path: Path, include_headers: bool = False, include_footers: bool = False) -> Path:
    lines: List[str] = []
    for page in document.get("pages", []):
        pnum = page.get("page_number", "?")
        lines.append(f"<!-- Pagina {pnum} -->")
        if include_headers and page.get("header"):
            lines.append(page["header"])
            lines.append("")
        for block in page.get("blocks", []):
            btype = block.get("type", "paragraph")
            txt = _block_text_md(block).strip()
            if btype == "heading":
                level = int(block.get("level", 2))
                level = max(1, min(3, level))
                lines.append("#" * level + " " + txt)
            elif btype == "list":
                ltype = block.get("list_type", "unordered")
                items = block.get("items", []) if isinstance(block.get("items"), list) else []
                if items:
                    for i, item in enumerate(items, start=1):
                        itxt = str(item.get("content", "")).strip()
                        lines.append((f"{i}. " if ltype == "ordered" else "- ") + itxt)
                else:
                    for i, txt_item in enumerate(_fallback_list_items_from_content(str(block.get("content", ""))), start=1):
                        lines.append((f"{i}. " if ltype == "ordered" else "- ") + txt_item)
            elif btype == "table":
                headers = block.get("headers", []) if isinstance(block.get("headers"), list) else []
                rows = block.get("rows", []) if isinstance(block.get("rows"), list) else []
                if headers:
                    lines.append("| " + " | ".join(map(str, headers)) + " |")
                    lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                for row in rows:
                    if isinstance(row, list):
                        lines.append("| " + " | ".join(map(str, row)) + " |")
            elif btype == "signature_block":
                lines.append("---")
                lines.append(txt)
            elif btype == "separator":
                lines.append("---")
            elif btype in {"image_placeholder", "stamp_or_seal"}:
                desc = block.get("description") or txt or "elemento non testuale"
                lines.append(f"[Nota: {desc}]")
            else:
                lines.append(txt)
            lines.append("")
        if include_footers and page.get("footer"):
            lines.append(page["footer"])
            lines.append("")
    out_path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return out_path
def export_html_faithful(input_path: Path, out_path: Path) -> Path:
    pages: List[str] = []
    suffix = input_path.suffix.lower()
    if suffix == ".pdf":
        doc = fitz.open(input_path)
        for i in range(len(doc)):
            pix = doc[i].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            data = base64.b64encode(pix.tobytes("png")).decode("utf-8")
            pages.append(f"<section class='page'><img src='data:image/png;base64,{data}' alt='pagina {i+1}'/></section>")
    else:
        mime = "image/jpeg" if suffix in {".jpg", ".jpeg"} else "image/png"
        data = base64.b64encode(input_path.read_bytes()).decode("utf-8")
        pages.append(f"<section class='page'><img src='data:{mime};base64,{data}' alt='pagina 1'/></section>")
    html = """<html><head><meta charset='utf-8'><style>
    body{background:#eee;margin:0;padding:24px 0;}
    .page{max-width:980px;margin:0 auto 20px auto;background:#fff;box-shadow:0 2px 10px rgba(0,0,0,.18)}
    .page img{width:100%;height:auto;display:block}
    </style></head><body>""" + "".join(pages) + "</body></html>"
    out_path.write_text(html, encoding="utf-8")
    return out_path
def _find_tesseract_exe() -> str | None:
    direct = shutil.which("tesseract")
    if direct:
        return direct
    candidates = [
        Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
        Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
    ]
    local = Path(os.environ.get("LOCALAPPDATA", ""))
    if local:
        candidates.extend(local.glob("**/Tesseract-OCR/tesseract.exe"))
    for c in candidates:
        if Path(c).exists():
            return str(c)
    return None
def export_html(document: Dict[str, Any], out_path: Path) -> Path:
    html_out: List[str] = ["<html><head><meta charset='utf-8'><style>body{font-family:Times New Roman,serif;} .page{margin:24px auto;max-width:900px;padding:24px;border:1px solid #ddd;} .sig{margin-top:18px;} .note{color:#666;font-style:italic;} p,li,h1,h2,h3{white-space:pre-wrap;}</style></head><body>"]
    for page in document.get("pages", []):
        html_out.append("<section class='page'>")
        html_out.append(f"<!-- Pagina {page.get('page_number')} -->")
        for block in page.get("blocks", []):
            t = block.get("type", "paragraph")
            content = str(block.get("content", ""))
            style = block.get("style", {}) if isinstance(block.get("style"), dict) else {}
            inline = _render_inline_html(block.get("inline_spans", []) if isinstance(block.get("inline_spans"), list) else [], content)
            css = _align_css(style)
            if t == "heading":
                lvl = max(1, min(3, int(block.get("level", 2))))
                html_out.append(f"<h{lvl} style='{css}'>{inline}</h{lvl}>")
            elif t == "list":
                items = block.get("items", []) if isinstance(block.get("items"), list) else []
                if not items:
                    items = [{"content": x} for x in _fallback_list_items_from_content(content)]
                tag = "ol" if block.get("list_type") == "ordered" else "ul"
                html_out.append(f"<{tag} style='{css}'>")
                for it in items:
                    html_out.append(f"<li>{html.escape(str(it.get('content','')))}</li>")
                html_out.append(f"</{tag}>")
            elif t == "table":
                html_out.append(f"<table border='1' cellspacing='0' cellpadding='4' style='{css}'>")
                headers = block.get("headers", []) if isinstance(block.get("headers"), list) else []
                if headers:
                    html_out.append("<tr>" + "".join([f"<th>{html.escape(str(h))}</th>" for h in headers]) + "</tr>")
                for r in block.get("rows", []) if isinstance(block.get("rows"), list) else []:
                    if isinstance(r, list):
                        html_out.append("<tr>" + "".join([f"<td>{html.escape(str(c))}</td>" for c in r]) + "</tr>")
                html_out.append("</table>")
            elif t == "signature_block":
                html_out.append(f"<div class='sig' style='{css}'><hr/><p>{inline}</p></div>")
            elif t == "separator":
                html_out.append("<hr/>")
            elif t in {"image_placeholder", "stamp_or_seal"}:
                html_out.append(f"<p class='note' style='{css}'>[{html.escape(str(block.get('description') or content))}]</p>")
            else:
                html_out.append(f"<p style='{css}'>{inline}</p>")
        html_out.append("</section>")
    html_out.append("</body></html>")
    out_path.write_text("\n".join(html_out), encoding="utf-8")
    return out_path
def export_docx(document: Dict[str, Any], out_path: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Pt(72)
    section.bottom_margin = Pt(72)
    section.left_margin = Pt(72)
    section.right_margin = Pt(72)
    def size_pt(sz: str) -> int:
        if sz == "small":
            return 9
        if sz == "large":
            return 14
        return 11
    for pidx, page in enumerate(document.get("pages", []), start=1):
        if pidx > 1:
            doc.add_page_break()
        for block in page.get("blocks", []):
            t = block.get("type", "paragraph")
            style = block.get("style", {}) if isinstance(block.get("style"), dict) else {}

            # list/table hanno contenuto strutturato (items/rows) e NON vogliono un paragrafo
            # iniziale col content grezzo: nel DOCX finale duplicava lista/tabella.
            if t == "list":
                items = block.get("items", []) if isinstance(block.get("items"), list) else []
                if not items:
                    items = [{"content": x} for x in _fallback_list_items_from_content(str(block.get("content", "")))]
                for item in items:
                    lp = doc.add_paragraph(style="List Number" if block.get("list_type") == "ordered" else "List Bullet")
                    lp.add_run(str(item.get("content", "")))
                continue

            if t == "table":
                rows = block.get("rows", []) if isinstance(block.get("rows"), list) else []
                headers = block.get("headers", []) if isinstance(block.get("headers"), list) else []
                cols = max(len(headers), len(rows[0]) if rows and isinstance(rows[0], list) else 1)
                table = doc.add_table(rows=1 if headers else 0, cols=cols)
                if headers:
                    for i, h in enumerate(headers):
                        table.rows[0].cells[i].text = str(h)
                for row in rows:
                    if isinstance(row, list):
                        cells = table.add_row().cells
                        for i, c in enumerate(row[:cols]):
                            cells[i].text = str(c)
                continue

            if t == "heading":
                level = max(1, min(3, int(block.get("level", 2))))
                p = doc.add_heading(level=level)
            else:
                p = doc.add_paragraph()
            align = str(style.get("alignment", "left")).lower()
            if align == "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif align == "right":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif align == "justify":
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            spans = block.get("inline_spans", []) if isinstance(block.get("inline_spans"), list) else []
            if not spans:
                spans = [{"text": block.get("content", "")}]
            for sp in spans:
                r = p.add_run(str(sp.get("text", "")))
                r.bold = bool(sp.get("bold", False))
                r.italic = bool(sp.get("italic", False))
                r.underline = bool(sp.get("underline", False))
                r.font.name = "Times New Roman"
                r.font.size = Pt(size_pt(style.get("font_size_relative", "normal")))
    doc.save(out_path)
    return out_path
def _pdf_looks_valid(path: Path) -> bool:
    try:
        if not path.exists() or path.stat().st_size < 1024:
            return False
        with path.open("rb") as fh:
            head = fh.read(5)
        return head == b"%PDF-"
    except OSError:
        return False


def export_searchable_pdf(input_path: Path, out_path: Path) -> Path:
    import importlib.util
    import logging
    import subprocess
    import sys
    if importlib.util.find_spec("ocrmypdf") is None:
        raise RuntimeError(
            "ocrmypdf non installato nell'ambiente corrente. "
            "Esegui setup_full_update.cmd o installa manualmente ocrmypdf + dipendenze sistema."
        )
    tess = _find_tesseract_exe()
    if not tess:
        raise RuntimeError("tesseract non trovato (né nel PATH né nei percorsi standard)")
    env = os.environ.copy()
    tess_dir = str(Path(tess).parent)
    env["PATH"] = tess_dir + os.pathsep + env.get("PATH", "")
    if out_path.exists():
        out_path.unlink()
    cmd = [
        sys.executable,
        "-m",
        "ocrmypdf",
        str(input_path),
        str(out_path),
        "--skip-text",
        "--output-type",
        "pdf",
    ]
    try:
        subprocess.check_call(cmd, env=env)
    except subprocess.CalledProcessError as exc:
        # ocrmypdf 16.10 + pikepdf>=9 fallisce in validazione finale con exit 15
        # (AttributeError: 'Pdf' object has no attribute 'check'), ma il PDF ricercabile
        # è già stato scritto su disco prima del crash. Se il file esiste ed è un PDF
        # valido, accettiamolo comunque e segnaliamo il problema nel log — così
        # l'utente ottiene il searchable.pdf invece di una conversione fallita.
        if _pdf_looks_valid(out_path):
            logging.warning(
                "ocrmypdf uscito con exit %s ma il file di output esiste ed è un PDF valido: "
                "probabilmente errore in validazione finale (incompatibilità pikepdf>=9). "
                "Considero il PDF ricercabile generato. Per eliminare il warning: pin pikepdf<9.",
                exc.returncode,
            )
            return out_path
        raise
    return out_path